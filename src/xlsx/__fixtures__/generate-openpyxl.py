"""Generates the openpyxl fixtures used by the xlsx read/write tests.

Run from this folder: python generate-openpyxl.py (openpyxl 3.1.x). Output files are small and
committed; rerun only when a test needs a new feature.
"""

import datetime
import io
import struct
import zlib

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


def tiny_png() -> bytes:
    """A 2x2 red PNG, built by hand so Pillow is the only extra dependency openpyxl needs."""
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 2 for _ in range(2))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def styles_and_values() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Amount"
    ws["C1"] = "When"
    ws["D1"] = "Flag"
    rows = [("Alpha", 10.5, datetime.date(2024, 1, 15), True), ("Beta", -3, datetime.date(2024, 2, 29), False), ("Gamma", 1234567.891, datetime.datetime(2024, 3, 1, 13, 30), True)]
    for index, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row=index, column=col, value=value)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=12)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="000000"), top=Side(style="thin"))
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "#,##0.00"
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "yyyy-mm-dd"
    ws["A6"] = "Line one\nLine two"
    ws["A6"].alignment = Alignment(wrap_text=True, indent=1, text_rotation=45)
    ws["A7"] = "  padded  "
    ws["B6"] = "=SUM(B2:B4)"
    ws["B7"] = '=IF(B6>0,"positive","negative")'
    ws["C6"] = "=_xlfn.XLOOKUP(\"Beta\",A2:A4,B2:B4)"
    ws["D6"] = "#N/A"
    ws["E2"] = "locked off"
    ws["E2"].protection = Protection(locked=False)
    ws["E3"].fill = PatternFill("darkGrid", fgColor="FF0000", bgColor="00FF00")
    ws["E4"].font = Font(italic=True, underline="double", strike=True, vertAlign="superscript", color="FF00B050")
    ws.merge_cells("A9:C10")
    ws["A9"] = "Merged"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14.5
    ws.column_dimensions["F"].hidden = True
    ws.row_dimensions[6].height = 42
    ws.row_dimensions[8].hidden = True
    ws.freeze_panes = "B2"
    ws.sheet_properties.tabColor = "FF0000"
    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = 42
    hidden.sheet_state = "hidden"
    other = wb.create_sheet("Other Sheet")
    other["A1"] = "=Data!B2*2"
    other["A2"] = "='Other Sheet'!A1+1"
    wb.save("openpyxl-styles.xlsx")


def features() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Features"
    for r in range(1, 11):
        ws.cell(row=r, column=1, value=r)
        ws.cell(row=r, column=2, value=r * r)
    red = DifferentialStyle(font=Font(color="9C0006"), fill=PatternFill(bgColor="FFC7CE"))
    ws.conditional_formatting.add("A1:A10", CellIsRule(operator="greaterThan", formula=["5"], font=Font(color="9C0006"), fill=PatternFill(bgColor="FFC7CE")))
    ws.conditional_formatting.add("B1:B10", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
    ws.conditional_formatting.add("A1:A10", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    ws.conditional_formatting.add("B1:B10", IconSetRule("3Arrows", "percent", [0, 33, 67]))
    ws.conditional_formatting.add("C1:C10", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=PatternFill(bgColor="DDEBF7")))
    text_rule = Rule(type="containsText", operator="containsText", text="err", dxf=red)
    text_rule.formula = ['NOT(ISERROR(SEARCH("err",D1)))']
    ws.conditional_formatting.add("D1:D10", text_rule)
    dv = DataValidation(type="list", formula1='"Red,Green,Blue"', allow_blank=True)
    dv.error = "Pick a colour"
    dv.errorTitle = "Invalid"
    dv.prompt = "Choose"
    dv.promptTitle = "Colour"
    ws.add_data_validation(dv)
    dv.add("E1:E10")
    whole = DataValidation(type="whole", operator="between", formula1="1", formula2="100", showErrorMessage=True)
    ws.add_data_validation(whole)
    whole.add("F1:F5")
    ws["G1"] = "Link"
    ws["G1"].hyperlink = "https://example.com/path?q=1"
    ws["G2"] = "Internal"
    ws["G2"].hyperlink = "#Features!A1"
    ws["A1"].comment = Comment("A note\nsecond line", "Tester")
    ws["B3"].comment = Comment("Another", "Reviewer")
    ws.auto_filter.ref = "A1:B10"
    ws.print_area = "A1:F10"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.oddHeader.center.text = "Report"
    ws.protection.sheet = True
    ws.protection.formatCells = False
    tsheet = wb.create_sheet("Table")
    tsheet.append(["Region", "Q1", "Q2"])
    for row in [["North", 10, 12], ["South", 7, 9], ["East", 4, 15]]:
        tsheet.append(row)
    table = Table(displayName="Sales", ref="A1:C4")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False)
    tsheet.add_table(table)
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Quarterly"
    data = Reference(tsheet, min_col=2, min_row=1, max_col=3, max_row=4)
    cats = Reference(tsheet, min_col=1, min_row=2, max_row=4)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    tsheet.add_chart(chart, "E2")
    line = LineChart()
    line.add_data(Reference(tsheet, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    tsheet.add_chart(line, "E20")
    pie = PieChart()
    pie.add_data(Reference(tsheet, min_col=3, min_row=1, max_row=4), titles_from_data=True)
    pie.set_categories(cats)
    tsheet.add_chart(pie, "M2")
    img = Image(io.BytesIO(tiny_png()))
    img.anchor = "A8"
    tsheet.add_image(img)
    wb.defined_names["TaxRate"] = __import__("openpyxl").workbook.defined_name.DefinedName("TaxRate", attr_text="Table!$B$2")
    wb.properties.title = "Fixture"
    wb.properties.creator = "openpyxl generator"
    wb.save("openpyxl-features.xlsx")


def date1904() -> None:
    wb = Workbook()
    wb.epoch = __import__("openpyxl").utils.datetime.CALENDAR_MAC_1904
    ws = wb.active
    ws["A1"] = datetime.date(2020, 6, 1)
    ws["A1"].number_format = "d-mmm-yyyy"
    ws["A2"] = 1
    wb.save("openpyxl-1904.xlsx")


if __name__ == "__main__":
    styles_and_values()
    features()
    date1904()
